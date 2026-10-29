#!/usr/bin/env python3
"""把 scrape.js 产出的 item_*.json 整理成 SKU 明细表。

用法：
    python3 report.py --dir /abs/path            # 生成 CSV + XLSX
    python3 report.py --dir /abs/path --no-xlsx  # 只要 CSV

依赖：openpyxl（缺失时自动退化为只出 CSV）
"""
import argparse
import csv
import glob
import json
import os
import re
from datetime import datetime

POLICY_KEYS = ["官方立减", "限时立减", "首单礼金", "店铺新客减", "红包",
               "淘金币", "满减", "百亿补贴", "赠品", "88VIP"]


def norm_sold(s):
    s = (s or "").replace("已售", "").strip()
    m = re.match(r"([\d.]+)\s*万", s)
    if m:
        return str(int(float(m.group(1)) * 10000))
    m = re.match(r"(\d+)", s)
    return m.group(1) if m else s


def parse_policy(item):
    """把「已享受 / 可再享」的文本行拆成结构化政策。"""
    lines = (item.get("enjoyed") or []) + (item.get("more") or [])
    text = " ".join(lines)
    p = {k: "" for k in POLICY_KEYS}

    m = re.search(r"官方立减\s*(\d+)%\s*省([\d.]+)元", text)
    if m:
        p["官方立减"] = f"官方立减{m.group(1)}%（省{m.group(2)}元）"
    elif re.search(r"官方立减\s*(\d+)%", text):
        p["官方立减"] = "官方立减{}%".format(re.search(r"官方立减\s*(\d+)%", text).group(1))

    m = re.search(r"限时立减\s*([\d.]+)元", text)
    if m:
        p["限时立减"] = f"限时立减{m.group(1)}元"
    else:
        m = re.search(r"直降\s*([\d.]+)元", text)
        if m:
            p["限时立减"] = f"直降{m.group(1)}元"

    if "首单礼金" in text:
        amt = re.search(r"[¥￥]\s*([\d.]+)\s*首单礼金", text)
        p["首单礼金"] = f"首单礼金{amt.group(1) + '元' if amt else ''}"

    m = re.search(r"店铺新客减\s*([\d.]+)元", text)
    if m:
        p["店铺新客减"] = f"店铺新客减{m.group(1)}元"

    m = re.search(r"([\d.]+)元红包", text)
    if m:
        p["红包"] = f"{m.group(1)}元红包（限时）"

    m = re.search(r"淘金币已抵\s*([\d.]+)元", text)
    if m:
        p["淘金币"] = f"淘金币抵{m.group(1)}元"

    m = re.search(r"满\s*(\d+)\s*减\s*(\d+)", text)
    if m:
        p["满减"] = f"满{m.group(1)}减{m.group(2)}"

    if "百亿补贴" in text:
        p["百亿补贴"] = "百亿补贴"
    if re.search(r"送赠品|买\d得\d|买\d送\d", text + " " + (item.get("title") or "")):
        p["赠品"] = "有赠品/买赠"
    if "88VIP" in text:
        p["88VIP"] = "88VIP可用"
    return p


def short_label(p):
    seg = []
    if p["官方立减"]:
        seg.append(re.sub(r"（.*?）", "", p["官方立减"]))
    if p["限时立减"]:
        seg.append(p["限时立减"].replace("限时立减", "限时降"))
    if p["满减"]:
        seg.append(p["满减"])
    if p["百亿补贴"]:
        seg.append(p["百亿补贴"])
    return "/".join(seg)


def build(items, outdir):
    rows, summary = [], []
    for it in items:
        pol = parse_policy(it)
        tags = [v for v in pol.values() if v]
        policy_text = " + ".join(tags) if tags else "无叠加优惠"
        short = short_label(pol) or "无叠加优惠"
        prices = [s["finalPrice"] for s in it.get("skus", []) if s.get("finalPrice")]

        summary.append({
            "商品ID": it.get("itemId", ""),
            "商品标题": it.get("title", ""),
            "店铺": it.get("shopName", ""),
            "链接": it.get("url", ""),
            "活动": it.get("activity", ""),
            "SKU数": len(it.get("skus", [])),
            "销量": norm_sold(it.get("sold", "")),
            "到手价区间": f"{min(prices):.2f} - {max(prices):.2f}" if prices else "-",
            "促销政策": policy_text,
            "政策简写": short,
            "抓取时间": it.get("capturedAt", "")[:19].replace("T", " "),
        })

        for s in it.get("skus", []):
            lp, fp = s.get("listPrice"), s.get("finalPrice")
            rows.append({
                "商品ID": it.get("itemId", ""),
                "商品标题": it.get("title", ""),
                "店铺": it.get("shopName", ""),
                "规格/SKU": s.get("spec") or "默认规格",
                "SKU ID": s.get("skuId", ""),
                "优惠前价(元)": lp,
                "到手价(元)": fp,
                "到手价口径": it.get("finalPriceLabel", "") or "页面优惠后价",
                "优惠力度": f"{round((1 - fp / lp) * 100, 1)}%" if lp and fp else "",
                "立减金额(元)": round(lp - fp, 2) if lp and fp else "",
                "促销政策": policy_text,
                "政策简写": short,
                "库存": s.get("stock"),
                "备注": "已售罄/无货" if (fp is None or s.get("stock") == 0) else "",
            })
    return rows, summary


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def write_xlsx(rows, summary, path, outdir):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    HEAD = PatternFill("solid", fgColor="1F4E5A")
    HEADF = Font(color="FFFFFF", bold=True, size=10)
    wb = Workbook()

    def sheet(ws, cols, data, widths, money=()):
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(1, c)
            cell.fill, cell.font = HEAD, HEADF
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for r in data:
            ws.append([r.get(k) if r.get(k) is not None else "" for k in cols])
            for c in money:
                ws.cell(ws.max_row, c).number_format = "0.00"
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    ws1 = wb.active
    ws1.title = "SKU明细"
    cols = list(rows[0].keys())
    sheet(ws1, cols, rows, [15, 44, 16, 30, 17, 12, 11, 14, 10, 12, 46, 22, 8, 12],
          money=(6, 7, 10))

    ws2 = wb.create_sheet("商品汇总")
    cols2 = list(summary[0].keys())
    sheet(ws2, cols2, summary, [15, 46, 16, 40, 8, 7, 10, 16, 52, 24, 18])

    ws3 = wb.create_sheet("口径说明")
    notes = [
        ["项目", "说明"],
        ["取数方式", "读取详情页服务端渲染数据（__ICE_APP_CONTEXT__ 中的 skuBase / skuCore），不调用签名接口"],
        ["优惠前价", "页面「优惠前」划线价，即 SKU 原价"],
        ["到手价", "页面优惠后价：天猫多为「平台加补后」，C 店多为「店铺优惠后」；口径名见「到手价口径」列"],
        ["口径差异", "「平台加补后」含平台补贴，「店铺优惠后」只减店铺优惠，两者不可直接横比，需看列内口径名"],
        ["优惠力度", "1 - 到手价 / 优惠前价，名义折扣率（非到手价占原价的实际支付比）"],
        ["促销政策", "详情页「已享受 / 可再享」逐行拆分；含官方立减%、限时立减（直降）、首单礼金、店铺新客减、红包、淘金币、跨店满减、百亿补贴、赠品"],
        ["注意1", "「已享受」明细按默认展示 SKU 渲染，金额为单品口径；不同 SKU 以各自到手价为准"],
        ["注意2", "红包、限时立减含倒计时，属限时权益，可能随时失效"],
        ["注意3", "88VIP、消费券、淘金币余额因人而异，最终以下单页为准"],
        ["数据目录", os.path.abspath(outdir)],
        ["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M")],
    ]
    for r in notes:
        ws3.append(r)
    for c in (1, 2):
        ws3.cell(1, c).fill, ws3.cell(1, c).font = HEAD, HEADF
    ws3.column_dimensions["A"].width = 14
    ws3.column_dimensions["B"].width = 104

    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True, help="scrape.js 的 OUT_DIR")
    ap.add_argument("--no-xlsx", action="store_true")
    args = ap.parse_args()

    files = sorted(glob.glob(os.path.join(args.dir, "item_*.json")))
    if not files:
        raise SystemExit(f"在 {args.dir} 下没有找到 item_*.json")

    items = [json.load(open(f, encoding="utf-8")) for f in files]
    rows, summary = build(items, args.dir)
    if not rows:
        raise SystemExit("没有可输出的 SKU 数据")

    stamp = datetime.now().strftime("%Y%m%d")
    csv_path = os.path.join(args.dir, f"SKU价格与优惠_{stamp}.csv")
    write_csv(rows, csv_path)
    print(f"CSV  → {csv_path}  ({len(rows)} 行 / {len(items)} 商品)")

    if not args.no_xlsx:
        try:
            xlsx_path = os.path.join(args.dir, f"SKU价格与优惠_{stamp}.xlsx")
            write_xlsx(rows, summary, xlsx_path, args.dir)
            print(f"XLSX → {xlsx_path}")
        except ImportError:
            print("未安装 openpyxl，跳过 XLSX（pip3 install openpyxl 后可生成）")

    print("\n商品概览：")
    for s in summary:
        print(f"  {s['商品ID']} | {s['SKU数']:>2}SKU | 到手 {s['到手价区间']:>15} | {s['政策简写']}")


if __name__ == "__main__":
    main()
